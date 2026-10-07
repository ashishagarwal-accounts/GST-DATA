"""GSTR-1 workbook: a summary sheet plus one sheet per section in the column order of the GST offline
tool's Excel import template (rows 1-3 summary, header on row 4, data from row 5)."""
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.gstr1 import Gstr1

HEADER_FILL = PatternFill("solid", start_color="DCE6F1")
BOLD = Font(bold=True)


def _d(value: date | None) -> str:
    return value.strftime("%d-%b-%Y") if value else ""


def _n(value) -> float | int | str:
    if isinstance(value, Decimal):
        return float(value)  # Excel cells are numeric; exact values were rounded to paise upstream
    return value if value is not None else ""


# sheet name -> (section key, [(header, row -> value)])
SHEETS = {
    "b2b,sez,de": ("b2b", [
        ("GSTIN/UIN of Recipient", lambda r: r["ctin"]), ("Receiver Name", lambda r: r["name"]),
        ("Invoice Number", lambda r: r["inum"]), ("Invoice date", lambda r: _d(r["idt"])),
        ("Invoice Value", lambda r: r["val"]), ("Place Of Supply", lambda r: r["pos"]),
        ("Reverse Charge", lambda r: r["rchrg"]), ("Applicable % of Tax Rate", lambda r: ""),
        ("Invoice Type", lambda r: r["inv_typ"]), ("E-Commerce GSTIN", lambda r: ""),
        ("Rate", lambda r: r["rt"]), ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "b2cl": ("b2cl", [
        ("Invoice Number", lambda r: r["inum"]), ("Invoice date", lambda r: _d(r["idt"])),
        ("Invoice Value", lambda r: r["val"]), ("Place Of Supply", lambda r: r["pos"]),
        ("Applicable % of Tax Rate", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]), ("E-Commerce GSTIN", lambda r: ""),
    ]),
    "b2cs": ("b2cs", [
        ("Type", lambda r: r["typ"]), ("Place Of Supply", lambda r: r["pos"]), ("Applicable % of Tax Rate", lambda r: ""),
        ("Rate", lambda r: r["rt"]), ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
        ("E-Commerce GSTIN", lambda r: ""),
    ]),
    "cdnr": ("cdnr", [
        ("GSTIN/UIN of Recipient", lambda r: r["ctin"]), ("Receiver Name", lambda r: r["name"]),
        ("Note Number", lambda r: r["nt_num"]), ("Note Date", lambda r: _d(r["nt_dt"])),
        ("Note Type", lambda r: r["ntty"]), ("Place Of Supply", lambda r: r["pos"]),
        ("Reverse Charge", lambda r: r["rchrg"]), ("Note Supply Type", lambda r: r["inv_typ"]),
        ("Note Value", lambda r: r["val"]), ("Applicable % of Tax Rate", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "cdnur": ("cdnur", [
        ("UR Type", lambda r: r["ur_typ"]), ("Note Number", lambda r: r["nt_num"]), ("Note Date", lambda r: _d(r["nt_dt"])),
        ("Note Type", lambda r: r["ntty"]), ("Place Of Supply", lambda r: r["pos"]), ("Note Value", lambda r: r["val"]),
        ("Applicable % of Tax Rate", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "exp": ("exp", [
        ("Export Type", lambda r: r["exp_typ"]), ("Invoice Number", lambda r: r["inum"]),
        ("Invoice date", lambda r: _d(r["idt"])), ("Invoice Value", lambda r: r["val"]), ("Port Code", lambda r: ""),
        ("Shipping Bill Number", lambda r: ""), ("Shipping Bill Date", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Taxable Value", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "at": ("at", [
        ("Place Of Supply", lambda r: r["pos"]), ("Applicable % of Tax Rate", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Gross Advance Received", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "atadj": ("atadj", [
        ("Place Of Supply", lambda r: r["pos"]), ("Applicable % of Tax Rate", lambda r: ""), ("Rate", lambda r: r["rt"]),
        ("Gross Advance Adjusted", lambda r: r["txval"]), ("Cess Amount", lambda r: r["csamt"]),
    ]),
    "exemp": ("exemp", [
        ("Description", lambda r: r["desc"]), ("Nil Rated Supplies", lambda r: r["nil"]),
        ("Exempted(other than nil rated/non GST supply)", lambda r: r["exempt"]),
        ("Non-GST Supplies", lambda r: r["non_gst"]),
    ]),
    **{name: (key, [
        ("HSN", lambda r: r["hsn"]), ("Description", lambda r: r["desc"]), ("UQC", lambda r: r["uqc"]),
        ("Total Quantity", lambda r: r["qty"]), ("Total Value", lambda r: r["val"]), ("Rate", lambda r: r["rt"]),
        ("Taxable Value", lambda r: r["txval"]), ("Integrated Tax Amount", lambda r: r["iamt"]),
        ("Central Tax Amount", lambda r: r["camt"]), ("State/UT Tax Amount", lambda r: r["samt"]),
        ("Cess Amount", lambda r: r["csamt"]),
    ]) for name, key in (("hsn(b2b)", "hsn_b2b"), ("hsn(b2c)", "hsn_b2c"))},
    "docs": ("docs", [
        ("Nature of Document", lambda r: r["nature"]), ("Sr. No. From", lambda r: r["from"]),
        ("Sr. No. To", lambda r: r["to"]), ("Total Number", lambda r: r["total"]), ("Cancelled", lambda r: r["cancelled"]),
    ]),
}


def build_workbook(report: Gstr1) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    g = report.gstin
    ws.append([f"GSTR-1 - {g.gstin} ({g.trade_name or g.legal_name or g.entity.name}) - period {report.period}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Prepared by the Alcove GST tool for review. Section sheets follow the GST offline tool template; "
               "verify by importing into the offline tool before filing."])
    ws.append([])
    header = ["Section", "GSTR-1 table", "Records", "Taxable value", "IGST", "CGST", "SGST", "Cess"]
    ws.append(header)
    for cell in ws[4]:
        cell.font, cell.fill = BOLD, HEADER_FILL
    for s in report.summary():
        ws.append([s["title"], s["table"], s["records"], *(_n(s[k]) for k in ("txval", "iamt", "camt", "samt", "csamt"))])
    liab = report.liability()
    ws.append([])
    ws.append(["Tax payable from this return (excl. reverse charge)", "", "", _n(liab["txval"]),
               *(_n(liab[k]) for k in ("iamt", "camt", "samt", "csamt"))])
    ws.cell(ws.max_row, 1).font = BOLD
    if report.warnings:
        ws.append([])
        ws.append(["Items to check"])
        ws.cell(ws.max_row, 1).font = BOLD
        for w in report.warnings:
            ws.append([w])
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16

    for name, (key, columns) in SHEETS.items():
        rows = report.sections[key]
        sh = wb.create_sheet(name)
        sh.append([f"Summary for {name}"])
        sh.append(["Records", len(rows)])
        sh.append([])
        sh.append([h for h, _ in columns])
        for cell in sh[4]:
            cell.font, cell.fill = BOLD, HEADER_FILL
            sh.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 2)
        for r in rows:
            sh.append([_n(f(r)) for _, f in columns])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
