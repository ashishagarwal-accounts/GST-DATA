"""Parse sales / purchase registers from Tally exports or hand-maintained Excel/CSV.

Each spreadsheet row is one invoice line (one HSN / tax-rate slab). Rows sharing an invoice key
are grouped into one invoice. Column headers are matched against an alias list, so the header
row may sit anywhere in the first 30 rows (Tally puts company name and period above it).

Errors block the import; warnings are kept on the batch for review.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import PurePath

from app.gstin import FOREIGN_STATE_CODE, gstin_error, normalize_gstin, parse_state_code
from app.ingest.values import (
    Issue, clean_text, is_blank, parse_bool, parse_date, parse_decimal, period_bounds,
)
from app.models import Direction, DocType, SupplyCategory

ZERO = Decimal("0.00")
ROUNDING_TOLERANCE = Decimal("1.00")
STANDARD_RATES = [Decimal(r) for r in ("0", "0.1", "0.25", "1", "1.5", "3", "5", "6", "7.5", "12", "18", "28", "40")]
HEADER_SCAN_ROWS = 30

COST_CENTRE_ALIASES = ("cost centre", "cost center", "cost centre name", "cost center name", "branch", "project", "site")

# canonical field -> accepted header spellings (after normalize_header)
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_no": ("invoice no", "invoice number", "voucher no", "vch no", "voucher number", "bill no",
                   "bill number", "doc no", "document no", "document number", "inv no", "supplier invoice no",
                   "note no", "note number"),
    "invoice_date": ("invoice date", "date", "voucher date", "vch date", "bill date", "doc date",
                     "document date", "inv date", "supplier invoice date", "note date"),
    "doc_type": ("doc type", "document type", "voucher type", "vch type", "note type"),
    "counterparty_gstin": ("gstin", "gstin/uin", "gstin uin", "party gstin", "party gstin/uin", "gstin of supplier",
                           "gstin of recipient", "gstin/uin of recipient", "supplier gstin", "customer gstin",
                           "gst no", "gst number", "gstin no"),
    "counterparty_name": ("party name", "party", "particulars", "customer name", "supplier name", "customer",
                          "supplier", "receiver name", "recipient name", "trade/legal name", "ledger name",
                          "party ledger name"),
    "place_of_supply": ("place of supply", "pos", "place of supply (state)"),
    "reverse_charge": ("reverse charge", "rcm", "reverse charge applicable", "rcm applicable", "is rcm"),
    "supply_category": ("supply category", "supply type", "category", "invoice type"),
    "itc_eligible": ("itc eligible", "eligible for itc", "itc eligibility", "itc available"),
    "hsn": ("hsn", "hsn/sac", "hsn code", "sac", "sac code", "hsn/sac code", "hsn sac"),
    "description": ("description", "item", "item name", "stock item", "item description", "goods/services"),
    "uqc": ("uqc", "uom", "unit"),
    "quantity": ("quantity", "qty"),
    "tax_rate": ("tax rate", "gst rate", "gst %", "gst rate %", "rate of tax", "rate (%)", "tax %", "gst rate (%)"),
    "taxable_value": ("taxable value", "taxable amount", "assessable value", "taxable val", "taxable"),
    "igst": ("igst", "igst amount", "integrated tax", "integrated tax amount", "igst amt"),
    "cgst": ("cgst", "cgst amount", "central tax", "central tax amount", "cgst amt"),
    "sgst": ("sgst", "sgst amount", "sgst/utgst", "sgst/utgst amount", "state tax", "state/ut tax",
             "state/ut tax amount", "utgst", "sgst amt"),
    "cess": ("cess", "cess amount", "cess amt"),
    "invoice_value": ("invoice value", "invoice amount", "gross total", "total value", "invoice total",
                      "total invoice value", "bill amount"),
    "service_type": ("service type", "type of service", "service category"),
    "cost_centre": COST_CENTRE_ALIASES,
}
REQUIRED_FIELDS = ("invoice_no", "invoice_date", "taxable_value")
_ALIAS_LOOKUP = {alias: canonical for canonical, aliases in COLUMN_ALIASES.items() for alias in aliases}
# Purchase registers carry both our voucher no/date and the supplier's; reconciliation needs the supplier's.
PREFERRED_ALIASES = {"supplier invoice no", "supplier invoice date"}


def normalize_header(value) -> str:
    text = str(value or "").lower().replace("_", " ").replace("\n", " ")
    text = re.sub(r"\((rs\.?|inr|₹|in rs\.?)\)", "", text)
    text = re.sub(r"\s+", " ", text).strip(" .:")
    return text


@dataclass
class ParsedItem:
    tax_rate: Decimal
    taxable_value: Decimal
    igst: Decimal = ZERO
    cgst: Decimal = ZERO
    sgst: Decimal = ZERO
    cess: Decimal = ZERO
    hsn: str | None = None
    description: str | None = None
    uqc: str | None = None
    quantity: Decimal | None = None
    service_type: str | None = None
    source_row: int | None = None
    land_value: Decimal = ZERO  # set at import for land-deduction service types
    cost_centre: str | None = None


@dataclass
class ParsedInvoice:
    doc_type: DocType
    invoice_no: str
    invoice_date: date
    counterparty_gstin: str | None
    counterparty_name: str | None
    place_of_supply: str | None
    reverse_charge: bool
    supply_category: SupplyCategory
    itc_eligible: bool | None
    invoice_value: Decimal
    items: list[ParsedItem]
    source_rows: list[int]

    def total(self, attr: str) -> Decimal:
        return sum((getattr(i, attr) for i in self.items), ZERO)


@dataclass
class RegisterParseResult:
    invoices: list[ParsedInvoice] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    sheet: str | None = None
    header_row: int | None = None
    column_map: dict[str, str] = field(default_factory=dict)
    unmapped_columns: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- file reading

def read_tables(content: bytes, filename: str) -> list[tuple[str, list[list]]]:
    """Return [(sheet_name, rows)] for xlsx/xlsm/xls/csv files."""
    suffix = PurePath(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        tables = []
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            tables.append((sheet.name, rows))
        return tables
    if suffix in (".csv", ".txt"):
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        return [(PurePath(filename).stem, [list(r) for r in csv.reader(io.StringIO(text))])]
    raise ValueError(f"unsupported file type '{suffix}' (use .xlsx, .xls or .csv)")


def find_header(rows: list[list], lookup: dict[str, str] | None = None,
                required: tuple[str, ...] = REQUIRED_FIELDS) -> tuple[int, dict[int, str], list[str]] | None:
    """Locate the header row. Returns (row_index, {col_index: canonical_field}, unmapped headers)."""
    lookup = lookup or _ALIAS_LOOKUP
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        preferred: set[str] = set()
        unmapped: list[str] = []
        for col, cell in enumerate(row):
            name = normalize_header(cell)
            if not name:
                continue
            canonical = lookup.get(name)
            taken_by = next((c for c, f in mapping.items() if f == canonical), None)
            if canonical and taken_by is None:
                mapping[col] = canonical
            elif canonical and name in PREFERRED_ALIASES and canonical not in preferred:
                unmapped.append(str(row[taken_by]).strip())
                del mapping[taken_by]
                mapping[col] = canonical
            else:
                unmapped.append(str(cell).strip())
                continue
            if name in PREFERRED_ALIASES:
                preferred.add(canonical)
        if all(f in mapping.values() for f in required):
            return idx, mapping, unmapped
    return None


# ---------------------------------------------------------------- field parsing

def parse_doc_type(value) -> DocType:
    text = (clean_text(value) or "").lower()
    if "credit" in text or text in ("c", "cn", "crn"):
        return DocType.CREDIT_NOTE
    if "debit" in text or text in ("d", "dn", "dbn"):
        return DocType.DEBIT_NOTE
    return DocType.INVOICE


def parse_supply_category(value) -> SupplyCategory:
    text = (clean_text(value) or "").lower().replace("-", " ")
    if not text or text in ("taxable", "regular", "b2b", "b2c", "b2cl", "b2cs", "normal", "r"):
        return SupplyCategory.TAXABLE
    without = any(w in text for w in ("without", "wopay", "lut", "bond"))
    if "deemed" in text:
        return SupplyCategory.DEEMED_EXPORT
    if "sez" in text:
        return SupplyCategory.SEZ_WITHOUT_PAYMENT if without else SupplyCategory.SEZ_WITH_PAYMENT
    if "export" in text or text in ("exp", "expwp", "expwop"):
        if without or text == "expwop":
            return SupplyCategory.EXPORT_WITHOUT_PAYMENT
        return SupplyCategory.EXPORT_WITH_PAYMENT
    if "nil" in text:
        return SupplyCategory.NIL_RATED
    if "exempt" in text:
        return SupplyCategory.EXEMPT
    if "non gst" in text or "nongst" in text:
        return SupplyCategory.NON_GST
    for member in SupplyCategory:
        if text == member.value.replace("_", " "):
            return member
    raise ValueError(f"unknown supply category '{value}'")


def infer_rate(taxable: Decimal, tax: Decimal) -> Decimal | None:
    if taxable == ZERO:
        return ZERO if tax == ZERO else None
    best = min(STANDARD_RATES, key=lambda r: abs(taxable * r / 100 - tax))
    tolerance = max(ROUNDING_TOLERANCE, tax * Decimal("0.005"))
    return best if abs(taxable * best / 100 - tax) <= tolerance else None


ZERO_TAX_CATEGORIES = {
    SupplyCategory.EXPORT_WITHOUT_PAYMENT, SupplyCategory.SEZ_WITHOUT_PAYMENT,
    SupplyCategory.NIL_RATED, SupplyCategory.EXEMPT, SupplyCategory.NON_GST,
}


# ---------------------------------------------------------------- main entry point

def parse_register(
    content: bytes, filename: str, *, direction: Direction, own_gstin: str, period: str,
) -> RegisterParseResult:
    result = RegisterParseResult()
    try:
        tables = read_tables(content, filename)
    except Exception as exc:  # corrupt workbook, wrong extension, etc.
        result.errors.append(Issue(f"could not read file: {exc}"))
        return result

    for sheet_name, rows in tables:
        found = find_header(rows)
        if found:
            header_idx, mapping, unmapped = found
            result.sheet, result.header_row = sheet_name, header_idx + 1
            result.column_map = {f: str(rows[header_idx][c]).strip() for c, f in mapping.items()}
            result.unmapped_columns = unmapped
            break
    else:
        result.errors.append(Issue(
            "no header row found: need columns for invoice number, invoice date and taxable value "
            f"in the first {HEADER_SCAN_ROWS} rows of some sheet"
        ))
        return result

    own_state = own_gstin[:2]
    groups: dict[tuple, list[tuple[int, dict]]] = {}
    for offset, raw in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        record = {f: (raw[c] if c < len(raw) else None) for c, f in mapping.items()}
        if is_blank(record.get("invoice_no")):
            continue  # blank lines, Tally sub-total / grand-total rows
        line = _parse_line(record, offset, direction, result)
        if line is None:
            continue
        key = (line["doc_type"], line["invoice_no"].upper(), line["counterparty_gstin"] or "",
               "" if line["counterparty_gstin"] or direction is Direction.OUTWARD
               else (line["counterparty_name"] or "").upper())
        groups.setdefault(key, []).append((offset, line))

    start, end = period_bounds(period)
    for lines in groups.values():
        invoice = _build_invoice(lines, direction, own_state, result)
        if invoice is None:
            continue
        if direction is Direction.OUTWARD and not (start <= invoice.invoice_date <= end):
            result.warnings.append(Issue(
                f"invoice {invoice.invoice_no} dated {invoice.invoice_date:%d-%m-%Y} is outside {period}",
                row=invoice.source_rows[0], field="invoice_date"))
        if invoice.invoice_date > end:
            result.errors.append(Issue(
                f"invoice {invoice.invoice_no} dated {invoice.invoice_date:%d-%m-%Y} is after the end of {period}",
                row=invoice.source_rows[0], field="invoice_date"))
            continue
        result.invoices.append(invoice)

    if not result.invoices and not result.errors:
        result.errors.append(Issue("no invoice rows found below the header"))
    return result


def _parse_line(record: dict, row: int, direction: Direction, result: RegisterParseResult) -> dict | None:
    errors_before = len(result.errors)

    def get(field_name, parser, default=None):
        try:
            value = parser(record.get(field_name))
        except ValueError as exc:
            result.errors.append(Issue(str(exc), row=row, field=field_name))
            return default
        return default if value is None else value

    line = {
        "invoice_no": clean_text(record.get("invoice_no"), 50),
        "invoice_date": get("invoice_date", parse_date),
        "doc_type": parse_doc_type(record.get("doc_type")),
        "counterparty_name": clean_text(record.get("counterparty_name"), 255),
        "reverse_charge": get("reverse_charge", parse_bool, False),
        "supply_category": get("supply_category", parse_supply_category, SupplyCategory.TAXABLE),
        "itc_eligible": get("itc_eligible", parse_bool, True) if direction is Direction.INWARD else None,
        "place_of_supply": get("place_of_supply", parse_state_code),
        "hsn": clean_text(record.get("hsn"), 8),
        "service_type": clean_text(record.get("service_type"), 100),
        "cost_centre": clean_text(record.get("cost_centre"), 100),
        "description": clean_text(record.get("description"), 255),
        "uqc": clean_text(record.get("uqc"), 10),
        "quantity": get("quantity", lambda v: parse_decimal(v, Decimal("0.001"))),
        "tax_rate": get("tax_rate", parse_decimal),
        "invoice_value": get("invoice_value", parse_decimal),
    }
    for amount in ("taxable_value", "igst", "cgst", "sgst", "cess"):
        line[amount] = get(amount, parse_decimal, ZERO)

    gstin_raw = clean_text(record.get("counterparty_gstin"))
    line["counterparty_gstin"] = None
    if gstin_raw and gstin_raw.upper() not in ("URP", "UNREGISTERED", "NA", "N/A", "-"):
        gstin = normalize_gstin(gstin_raw)
        problem = gstin_error(gstin)
        if problem:
            result.errors.append(Issue(problem, row=row, field="counterparty_gstin"))
        line["counterparty_gstin"] = gstin

    if line["invoice_date"] is None and len(result.errors) == errors_before:
        result.errors.append(Issue("invoice date is missing", row=row, field="invoice_date"))
    if len(result.errors) > errors_before:
        return None

    # Credit notes are often exported with negative amounts; store magnitudes, sign comes from doc_type.
    amounts = ("taxable_value", "igst", "cgst", "sgst", "cess", "invoice_value")
    if any((line[a] or ZERO) < 0 for a in amounts):
        if line["doc_type"] is DocType.INVOICE:
            result.errors.append(Issue(
                "negative amounts on an invoice row; mark it as a credit note in the document type column",
                row=row, field="taxable_value"))
            return None
        for a in amounts:
            if line[a] is not None:
                line[a] = abs(line[a])

    if line["igst"] and (line["cgst"] or line["sgst"]):
        result.errors.append(Issue("row has both IGST and CGST/SGST", row=row, field="igst"))
        return None
    if abs(line["cgst"] - line["sgst"]) > ROUNDING_TOLERANCE:
        result.errors.append(Issue(f"CGST {line['cgst']} and SGST {line['sgst']} differ", row=row, field="cgst"))
        return None

    tax = line["igst"] + line["cgst"] + line["sgst"]
    if line["tax_rate"] is None:
        rate = infer_rate(line["taxable_value"], tax)
        if rate is None:
            result.errors.append(Issue(
                f"cannot infer tax rate from taxable {line['taxable_value']} and tax {tax}; add a tax rate column",
                row=row, field="tax_rate"))
            return None
        line["tax_rate"] = rate
    else:
        expected = line["taxable_value"] * line["tax_rate"] / 100
        if tax == ZERO and line["tax_rate"] > 0 and not line["reverse_charge"] \
                and line["supply_category"] not in ZERO_TAX_CATEGORIES:
            result.errors.append(Issue(f"tax rate {line['tax_rate']}% but no tax amounts", row=row, field="tax_rate"))
            return None
        if tax and abs(expected - tax) > max(ROUNDING_TOLERANCE, tax * Decimal("0.005")):
            result.warnings.append(Issue(
                f"tax {tax} does not match {line['tax_rate']}% of {line['taxable_value']} (expected {expected:.2f})",
                row=row, field="tax_rate"))
    return line


_INVOICE_LEVEL = ("invoice_date", "counterparty_gstin", "place_of_supply", "reverse_charge", "supply_category")


def _build_invoice(lines: list[tuple[int, dict]], direction: Direction, own_state: str,
                   result: RegisterParseResult) -> ParsedInvoice | None:
    rows = [r for r, _ in lines]
    first = lines[0][1]
    for attr in _INVOICE_LEVEL:
        values = {line[attr] for _, line in lines if line[attr] is not None}
        if len(values) > 1:
            result.errors.append(Issue(
                f"invoice {first['invoice_no']} has conflicting {attr.replace('_', ' ')} across rows {rows}",
                row=rows[0], field=attr))
            return None

    def first_set(attr):
        return next((line[attr] for _, line in lines if line[attr] is not None), None)

    items = [ParsedItem(
        tax_rate=line["tax_rate"], taxable_value=line["taxable_value"], igst=line["igst"], cgst=line["cgst"],
        sgst=line["sgst"], cess=line["cess"], hsn=line["hsn"], description=line["description"],
        uqc=line["uqc"], quantity=line["quantity"], service_type=line["service_type"], source_row=row,
        cost_centre=line["cost_centre"],
    ) for row, line in lines]

    gstin = first_set("counterparty_gstin")
    category = first_set("supply_category") or SupplyCategory.TAXABLE
    pos = first_set("place_of_supply")
    if pos is None:
        if category in (SupplyCategory.EXPORT_WITH_PAYMENT, SupplyCategory.EXPORT_WITHOUT_PAYMENT):
            pos = FOREIGN_STATE_CODE
        elif gstin:
            pos = gstin[:2]
        elif direction is Direction.OUTWARD:
            pos = own_state
            result.warnings.append(Issue(
                f"invoice {first['invoice_no']}: no place of supply or customer GSTIN, assumed own state {own_state}",
                row=rows[0], field="place_of_supply"))

    invoice = ParsedInvoice(
        doc_type=first["doc_type"], invoice_no=first["invoice_no"], invoice_date=first["invoice_date"],
        counterparty_gstin=gstin, counterparty_name=first_set("counterparty_name"), place_of_supply=pos,
        reverse_charge=bool(first_set("reverse_charge")), supply_category=category,
        itc_eligible=first_set("itc_eligible") if direction is Direction.INWARD else None,
        invoice_value=ZERO, items=items, source_rows=rows,
    )

    computed = sum((invoice.total(a) for a in ("taxable_value", "igst", "cgst", "sgst", "cess")), ZERO)
    stated = {line["invoice_value"] for _, line in lines if line["invoice_value"] is not None}
    if len(stated) == 1:
        invoice.invoice_value = stated.pop()
        if abs(invoice.invoice_value - computed) > ROUNDING_TOLERANCE:
            result.warnings.append(Issue(
                f"invoice {invoice.invoice_no}: stated value {invoice.invoice_value} differs from "
                f"taxable + tax {computed}", row=rows[0], field="invoice_value"))
    else:
        if len(stated) > 1:
            result.warnings.append(Issue(
                f"invoice {invoice.invoice_no}: rows carry different invoice values; using taxable + tax",
                row=rows[0], field="invoice_value"))
        invoice.invoice_value = computed

    if pos and pos != FOREIGN_STATE_CODE and category is SupplyCategory.TAXABLE:
        interstate = pos != own_state
        if direction is Direction.OUTWARD:
            if interstate and (invoice.total("cgst") or invoice.total("sgst")):
                result.warnings.append(Issue(
                    f"invoice {invoice.invoice_no}: place of supply {pos} is another state but CGST/SGST charged",
                    row=rows[0], field="place_of_supply"))
            if not interstate and invoice.total("igst"):
                result.warnings.append(Issue(
                    f"invoice {invoice.invoice_no}: place of supply is own state but IGST charged",
                    row=rows[0], field="place_of_supply"))

    if direction is Direction.INWARD and not gstin and not invoice.reverse_charge:
        result.warnings.append(Issue(
            f"purchase {invoice.invoice_no} has no supplier GSTIN and is not marked reverse charge",
            row=rows[0], field="counterparty_gstin"))
    return invoice
