import io
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_session
from app.models import AdvanceRule, Gstin, ServiceType
from app.routers.imports import _gstin_or_404, _period
from app.schemas import (
    AdvanceRuleCreate, AdvanceRuleRead, ServiceTypeCreate, ServiceTypeRead, ServiceTypeUpdate,
)
from app.services.advance_ledger import ledger_rows, period_report, run_ledger

router = APIRouter(tags=["advances"])


# ---------------------------------------------------------------- service types

def _gstin_by_id(session: Session, gstin_id: int) -> Gstin:
    gstin = session.get(Gstin, gstin_id)
    if not gstin:
        raise HTTPException(404, f"GSTIN record {gstin_id} not found")
    return gstin


@router.get("/gstins/{gstin_id}/service-types", response_model=list[ServiceTypeRead])
def list_service_types(gstin_id: int, session: Session = Depends(get_session)):
    _gstin_by_id(session, gstin_id)
    return session.scalars(select(ServiceType).where(ServiceType.gstin_id == gstin_id).order_by(ServiceType.name)).all()


@router.post("/gstins/{gstin_id}/service-types", response_model=ServiceTypeRead, status_code=201)
def create_service_type(gstin_id: int, body: ServiceTypeCreate, session: Session = Depends(get_session)):
    _gstin_by_id(session, gstin_id)
    clash = session.scalars(select(ServiceType).where(ServiceType.gstin_id == gstin_id)).all()
    if any(s.name.lower() == body.name.lower() for s in clash):
        raise HTTPException(409, f"service type '{body.name}' already exists for this GSTIN")
    service = ServiceType(gstin_id=gstin_id, **body.model_dump())
    session.add(service)
    session.commit()
    return service


@router.patch("/service-types/{service_type_id}", response_model=ServiceTypeRead)
def update_service_type(service_type_id: int, body: ServiceTypeUpdate, session: Session = Depends(get_session)):
    service = session.get(ServiceType, service_type_id)
    if not service:
        raise HTTPException(404, f"service type {service_type_id} not found")
    for key, value in body.model_dump(exclude_unset=True).items():
        setattr(service, key, value)
    session.commit()
    return service


# ---------------------------------------------------------------- ledger & GSTR-1 tables

def exact(value):
    """Money as exact decimal strings (like the Pydantic-backed endpoints), never floats."""
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {k: exact(v) for k, v in value.items()}
    if isinstance(value, list):
        return [exact(v) for v in value]
    return value


@router.get("/advances/report")
def advance_report(gstin: str, period: str, session: Session = Depends(get_session)):
    """GSTR-1 Table 11A / 11B for the period, plus movement and closing balance."""
    return exact(period_report(run_ledger(session, _gstin_or_404(session, gstin)), _period(period)))


@router.get("/advances/ledger")
def advance_ledger(gstin: str, period: str | None = None, session: Session = Depends(get_session)):
    """Every advance with its adjustments and refunds (up to the period, if given)."""
    record = _gstin_or_404(session, gstin)
    return exact(ledger_rows(run_ledger(session, record), _period(period) if period else None))


@router.get("/advances/rules", response_model=list[AdvanceRuleRead])
def list_rules(gstin: str, session: Session = Depends(get_session)):
    record = _gstin_or_404(session, gstin)
    return session.scalars(select(AdvanceRule).where(AdvanceRule.gstin_id == record.id).order_by(AdvanceRule.id)).all()


@router.post("/advances/rules", response_model=AdvanceRuleRead, status_code=201)
def create_rule(body: AdvanceRuleCreate, session: Session = Depends(get_session)):
    record = _gstin_or_404(session, body.gstin)
    rule = AdvanceRule(gstin_id=record.id, invoice_no=body.invoice_no.strip(),
                       advance_voucher_no=body.advance_voucher_no.strip(), amount=body.amount, note=body.note)
    session.add(rule)
    session.commit()
    return rule


@router.delete("/advances/rules/{rule_id}", status_code=204)
def delete_rule(rule_id: int, session: Session = Depends(get_session)):
    rule = session.get(AdvanceRule, rule_id)
    if not rule:
        raise HTTPException(404, f"rule {rule_id} not found")
    session.delete(rule)
    session.commit()


# ---------------------------------------------------------------- upload templates

TEMPLATES = {
    "advance_register": (
        ["Date", "Voucher No", "Type", "Customer Name", "Customer GSTIN", "Service Type", "Cost Centre",
         "Amount (incl. GST)", "Place of Supply", "Against Voucher No", "Narration"],
        [["05-09-2026", "RCPT/101", "Advance", "Rahul Sharma", "", "Flat Sale", "Tower A", 1000000, "", "", "Booking amount, Flat 3B"],
         ["12-09-2026", "RCPT/102", "Advance", "Tenant A Retail Pvt Ltd", "", "CAM", "Mall Road", 59000, "", "", "CAM advance Q3"],
         ["20-09-2026", "PYMT/7", "Refund", "Rahul Sharma", "", "Flat Sale", "Tower A", 200000, "", "RCPT/101", "Partial cancellation"],
         ["21-09-2026", "RCPT/103", "Other", "Bank interest", "", "", "", 1250, "", "", "Not an advance - skipped"]],
        "Type: Advance / Refund / Other (Other rows are skipped). Amount includes GST. "
        "Service Type must match one set up on the Advances page; Cost Centre (optional) one set up under Entities. Place of Supply defaults to the GSTIN's state.",
    ),
    "advance_opening": (
        ["Original Receipt Date", "Original Receipt No", "Customer Name", "Customer GSTIN", "Service Type",
         "Cost Centre", "Unadjusted Amount (incl. GST)", "Place of Supply", "Taxable Value", "CGST", "SGST", "IGST"],
        [["15-01-2026", "RCPT/0457", "Priya Das", "", "Flat Sale", "Tower A", 500000, "", "", "", "", ""]],
        "Unadjusted advances as at go-live, dated before the go-live month. Taxable Value / tax columns are "
        "optional: fill them with the tax actually paid if it differs from a back-calculation.",
    ),
    "sales_register": (
        ["Invoice No", "Invoice Date", "Doc Type", "Customer Name", "Customer GSTIN", "Place of Supply", "HSN/SAC",
         "Service Type", "Cost Centre", "Tax Rate", "Taxable Value", "IGST", "CGST", "SGST", "Invoice Value"],
        # 8,00,000 incl. GST at 18% with land deduction: consideration 7,14,285.71, tax 85,714.29 (18% of two-thirds)
        [["S/26-27/101", "30-09-2026", "Invoice", "Rahul Sharma", "", "", "995411", "Flat Sale", "Tower A", 18,
          714285.71, 0, 42857.14, 42857.15, 800000]],
        "One row per invoice line. Service Type links the line to advances of the same service. For land-deduction "
        "services enter the full consideration as Taxable Value; the tool reports two-thirds of it in GSTR-1.",
    ),
}


@router.get("/templates/{name}.xlsx")
def download_template(name: str):
    if name not in TEMPLATES:
        raise HTTPException(404, f"no template '{name}'")
    header, examples, note = TEMPLATES[name]
    wb = Workbook()
    ws = wb.active
    ws.title = "Upload"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 4)
    for row in examples:
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes.append([note])
    notes.append(["Delete the example rows before uploading. The header row may sit below title rows."])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}_template.xlsx"'})
