import io
from decimal import Decimal

from openpyxl import load_workbook

from tests.conftest import CUSTOMER_MH, CUSTOMER_WB, xlsx

D = Decimal
HEADER = ["Invoice No", "Invoice Date", "Doc Type", "Customer Name", "Customer GSTIN", "Place of Supply", "HSN/SAC",
          "Service Type", "Supply Type", "Tax Rate", "Taxable Value", "IGST", "CGST", "SGST"]
ROWS = [
    ["S/1", "01-09-2026", "Invoice", "Tenant WB", CUSTOMER_WB, None, "997212", None, None, 18, 100000, 0, 9000, 9000],
    ["S/2", "02-09-2026", "Invoice", "Tenant MH", CUSTOMER_MH, None, "997212", None, None, 18, 50000, 9000, 0, 0],
    # B2C flat sale: register has the full consideration, two-thirds is reported
    ["S/3", "03-09-2026", "Invoice", "Walk-in Buyer", None, None, "995411", "Flat Sale", None, 18, 89285.71, 0, 5357.14, 5357.15],
    ["CN/1", "04-09-2026", "Credit Note", "Walk-in Buyer", None, None, "995411", None, None, 18, 10000, 0, 900, 900],
    ["CN/2", "05-09-2026", "Credit Note", "Tenant WB", CUSTOMER_WB, None, "997212", None, None, 18, 20000, 0, 1800, 1800],
    ["S/4", "06-09-2026", "Invoice", "Mumbai Buyer", None, "27-Maharashtra", "995411", None, None, 18, 200000, 36000, 0, 0],
    ["S/5", "07-09-2026", "Invoice", "Residential Tenant", None, None, "997211", None, "Exempt", 0, 30000, 0, 0, 0],
    ["S/6", "08-09-2026", "Invoice", "Overseas Client", None, None, "998399", None, "Export without payment", 18, 100000, 0, 0, 0],
]


def setup(client, own_gstin):
    (entity,) = client.get("/api/entities").json()
    gid = entity["gstins"][0]["id"]
    client.post(f"/api/gstins/{gid}/service-types", json={"name": "Flat Sale", "tax_rate": 18, "land_deduction": True})
    r = client.post("/api/imports/register", files={"file": ("s.xlsx", xlsx([HEADER, *ROWS]))},
                    data={"gstin": own_gstin, "period": "2026-09", "kind": "sales_register"})
    assert r.status_code == 201, r.text
    r = client.get(f"/api/gstr1?gstin={own_gstin}&period=2026-09")
    assert r.status_code == 200, r.text
    return r.json()


def test_gstr1_classification(client, own_gstin):
    g = setup(client, own_gstin)
    s = g["sections"]
    assert [(r["inum"], r["inv_typ"]) for r in s["b2b"]] == [("S/1", "Regular B2B"), ("S/2", "Regular B2B")]
    assert [r["inum"] for r in s["b2cl"]] == ["S/4"]
    # B2C credit note CN/1 is netted off against the B2CS flat sale (59,523.81 two-thirds - 10,000)
    (b2cs,) = s["b2cs"]
    assert (b2cs["pos"], b2cs["rt"], b2cs["txval"]) == ("19-West Bengal", "18.00", "49523.81")
    assert (b2cs["camt"], b2cs["samt"]) == ("4457.14", "4457.15")
    # B2B credit note in its own table
    assert [(r["nt_num"], r["ntty"], r["ctin"]) for r in s["cdnr"]] == [("CN/2", "C", CUSTOMER_WB)]
    assert s["cdnur"] == []
    assert [(r["exp_typ"], r["inum"], r["iamt"]) for r in s["exp"]] == [("WOPAY", "S/6", "0.00")]
    assert s["exemp"] == [{"desc": "Intra-State supplies to unregistered persons", "nil": "0.00",
                           "exempt": "30000.00", "non_gst": "0.00"}]
    docs = {d["nature"]: d for d in s["docs"]}
    assert (docs["Invoices for outward supply"]["from"], docs["Invoices for outward supply"]["to"],
            docs["Invoices for outward supply"]["total"]) == ("S/1", "S/6", 6)
    assert docs["Credit Note"]["total"] == 2
    hsn = {(r["hsn"], r["uqc"]): r for r in s["hsn_b2c"]}
    assert hsn[("995411", "NA")]["txval"] == str(D("59523.81") - D("10000") + D("200000"))
    # 27,000 B2B + 36,000 B2CL + 8,914.29 net B2CS - 3,600 CDNR
    assert D(g["liability"]["tax"]) == D("68314.29")


def test_b2cs_negative_and_export_notes(client, own_gstin):
    g = setup(client, own_gstin)
    assert g["warnings"] == []
    r = client.post("/api/imports/register", files={"file": ("s.xlsx", xlsx([HEADER,
        ["CN/9", "10-10-2026", "Credit Note", "Walk-in Buyer", None, None, "995411", None, None, 18, 5000, 0, 450, 450],
        ["CN/10", "11-10-2026", "Credit Note", "Overseas Client", None, None, "998399", None, "Export without payment", 18, 1000, 0, 0, 0],
    ]))}, data={"gstin": own_gstin, "period": "2026-10", "kind": "sales_register"})
    assert r.status_code == 201, r.text
    oct_ = client.get(f"/api/gstr1?gstin={own_gstin}&period=2026-10").json()
    assert any("credit notes exceed sales" in w for w in oct_["warnings"])
    assert [(r["ur_typ"], r["nt_num"]) for r in oct_["sections"]["cdnur"]] == [("EXPWOP", "CN/10")]


def test_gstr1_excel_layout(client, own_gstin):
    setup(client, own_gstin)
    r = client.get(f"/api/gstr1/export.xlsx?gstin={own_gstin}&period=2026-09")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:4] == ["Summary", "b2b,sez,de", "b2cl", "b2cs"]
    b2b = wb["b2b,sez,de"]
    assert b2b["A4"].value == "GSTIN/UIN of Recipient" and b2b["A5"].value == CUSTOMER_WB
    assert b2b["D5"].value == "01-Sep-2026" and b2b["F5"].value == "19-West Bengal"
    assert wb["at"]["D4"].value == "Gross Advance Received"


def test_gstr1_portal_json(client, own_gstin):
    setup(client, own_gstin)
    r = client.get(f"/api/gstr1/export.json?gstin={own_gstin}&period=2026-09")
    assert r.status_code == 200
    assert r.headers["content-disposition"].endswith(f'GSTR1_{own_gstin}_092026.json"')
    j = r.json()
    assert (j["gstin"], j["fp"]) == (own_gstin, "092026")

    b2b = {x["ctin"]: x["inv"] for x in j["b2b"]}
    (s1,) = b2b[CUSTOMER_WB]
    assert (s1["inum"], s1["idt"], s1["pos"], s1["rchrg"], s1["inv_typ"], s1["val"]) == ("S/1", "01-09-2026", "19", "N", "R", 118000.0)
    assert s1["itms"] == [{"num": 1801, "itm_det": {"txval": 100000.0, "rt": 18, "iamt": 0.0, "camt": 9000.0, "samt": 9000.0, "csamt": 0.0}}]
    assert b2b[CUSTOMER_MH][0]["itms"][0]["itm_det"]["iamt"] == 9000.0

    assert j["b2cs"] == [{"sply_ty": "INTRA", "pos": "19", "typ": "OE", "rt": 18, "txval": 49523.81,
                          "iamt": 0.0, "camt": 4457.14, "samt": 4457.15, "csamt": 0.0}]
    assert [(x["pos"], x["inv"][0]["inum"]) for x in j["b2cl"]] == [("27", "S/4")]
    (cn,) = j["cdnr"][0]["nt"]
    assert (j["cdnr"][0]["ctin"], cn["ntty"], cn["nt_num"], cn["nt_dt"]) == (CUSTOMER_WB, "C", "CN/2", "05-09-2026")
    assert j["exp"] == [{"exp_typ": "WOPAY", "inv": [{"inum": "S/6", "idt": "08-09-2026", "val": 100000.0,
                         "itms": [{"txval": 100000.0, "rt": 18, "iamt": 0.0, "csamt": 0.0}]}]}]
    assert j["nil"] == {"inv": [{"sply_ty": "INTRAB2C", "nil_amt": 0.0, "expt_amt": 30000.0, "ngsup_amt": 0.0}]}
    # B2B: 997212 @18% (S/1, S/2, CN/2); B2C: 995411 @18%, 997211 @0% (exempt), 998399 @18% (export)
    assert {k: [(h["num"], h["hsn_sc"]) for h in v] for k, v in j["hsn"].items()} == {
        "hsn_b2b": [(1, "997212")], "hsn_b2c": [(1, "995411"), (2, "997211"), (3, "998399")]}
    docs = {d["doc_num"]: d["docs"][0] for d in j["doc_issue"]["doc_det"]}
    assert (docs[1]["from"], docs[1]["to"], docs[1]["totnum"], docs[1]["net_issue"]) == ("S/1", "S/6", 6, 6)
    assert docs[5]["totnum"] == 2
    assert "at" not in j and "txpd" not in j  # advances not enabled for this GSTIN
