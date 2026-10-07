import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_session
from app.models import CostCentre, Gstin
from app.routers.advances import exact
from app.routers.imports import _gstin_or_404, period_range
from app.schemas import CostCentreCreate, CostCentreRead, CostCentreUpdate
from app.services.cost_centres import cost_centre_report

router = APIRouter(tags=["cost centres"])


@router.get("/gstins/{gstin_id}/cost-centres", response_model=list[CostCentreRead])
def list_cost_centres(gstin_id: int, session: Session = Depends(get_session)):
    if not session.get(Gstin, gstin_id):
        raise HTTPException(404, f"GSTIN record {gstin_id} not found")
    return session.scalars(select(CostCentre).where(CostCentre.gstin_id == gstin_id).order_by(CostCentre.name)).all()


def _check_unique(session: Session, gstin_id: int, name: str, exclude_id: int | None = None) -> None:
    for c in session.scalars(select(CostCentre).where(CostCentre.gstin_id == gstin_id)):
        if c.id != exclude_id and c.name.lower() == name.lower():
            raise HTTPException(409, f"cost centre '{name}' already exists for this GSTIN")


@router.post("/gstins/{gstin_id}/cost-centres", response_model=CostCentreRead, status_code=201)
def create_cost_centre(gstin_id: int, body: CostCentreCreate, session: Session = Depends(get_session)):
    if not session.get(Gstin, gstin_id):
        raise HTTPException(404, f"GSTIN record {gstin_id} not found")
    _check_unique(session, gstin_id, body.name)
    centre = CostCentre(gstin_id=gstin_id, **body.model_dump())
    session.add(centre)
    session.commit()
    return centre


@router.patch("/cost-centres/{cost_centre_id}", response_model=CostCentreRead)
def update_cost_centre(cost_centre_id: int, body: CostCentreUpdate, session: Session = Depends(get_session)):
    centre = session.get(CostCentre, cost_centre_id)
    if not centre:
        raise HTTPException(404, f"cost centre {cost_centre_id} not found")
    changes = body.model_dump(exclude_unset=True)
    if changes.get("name"):
        changes["name"] = " ".join(changes["name"].split())
        _check_unique(session, centre.gstin_id, changes["name"], exclude_id=centre.id)
    for key, value in changes.items():
        setattr(centre, key, value)
    session.commit()
    return centre


@router.get("/reports/cost-centres")
def gst_by_cost_centre(gstin: str, months: list[str] = Depends(period_range), session: Session = Depends(get_session)):
    """Output tax, ITC and advance tax per cost centre for the period; rows add up to the GSTIN."""
    return exact(cost_centre_report(session, _gstin_or_404(session, gstin), months))


@router.get("/reports/cost-centres.xlsx")
def gst_by_cost_centre_excel(gstin: str, months: list[str] = Depends(period_range), session: Session = Depends(get_session)):
    record = _gstin_or_404(session, gstin)
    rep = cost_centre_report(session, record, months)
    span = rep["from"] if rep["from"] == rep["to"] else f"{rep['from']} to {rep['to']}"
    wb = Workbook()
    ws = wb.active
    ws.title = "GST by Cost Centre"
    ws.append([f"GST by Cost Centre - {record.gstin} ({record.entity.name}) - {span}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Indicative split; ITC cross-utilisation and the return are at GSTIN level."])
    ws.append([])
    header = ["Cost centre", "Code", "Sales docs", "Outward taxable", "Output tax", "Advance tax 11A",
              "Advance set-off 11B", "Purchase docs", "Inward taxable", "ITC (books)", "Net GST"]
    ws.append(header)
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", start_color="DCE6F1")
    keys = ("outward_taxable", "output_tax", "advance_11a", "advance_11b")
    for r in rep["rows"]:
        ws.append([r["name"], r["code"] or "", r["outward_docs"], *(float(r[k]) for k in keys),
                   r["inward_docs"], float(r["inward_taxable"]), float(r["itc"]), float(r["net"])])
    t = rep["totals"]
    ws.append(["Total", "", "", *(float(t[k]) for k in keys), "", float(t["inward_taxable"]), float(t["itc"]), float(t["net"])])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 17
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="GST_by_cost_centre_{record.gstin}_{span.replace(' to ', '_')}.xlsx"'})
